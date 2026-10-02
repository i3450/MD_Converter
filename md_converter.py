"""
md_converter.py: convierte PDF, Word, Excel, PowerPoint, HTML, CSV e imagenes
a Markdown para pasarselos a una IA gastando menos tokens.

Uso:
    python md_converter.py                # convierte los archivos de la carpeta actual
    python md_converter.py -i docs -r     # carpeta "docs", con subcarpetas
    python md_converter.py --force        # reconvierte todo
    python md_converter.py --benchmark    # compara los tokens contra MarkItDown puro
    python md_converter.py --help         # ver todas las opciones

Las imagenes se describen con OpenAI o Azure OpenAI si hay credenciales en un archivo .env
(opcional). Instalacion, configuracion y todas las opciones: ver https://github.com/i3450/MD_Converter.
"""

import argparse
import base64
import hashlib
import io
import json
import os
import re
import sys
import time
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from markitdown import MarkItDown
from openai import AzureOpenAI, OpenAI

try:
    import pdfplumber
except ImportError:
    pdfplumber = None

try:
    import openpyxl
    from openpyxl.utils import get_column_letter
except ImportError:
    openpyxl = None

try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None

try:
    from dotenv import load_dotenv  # pip install python-dotenv (opcional)
    load_dotenv()
except ImportError:
    pass

EXTENSIONS = {".pdf", ".docx", ".xlsx", ".pptx", ".html", ".csv", ".png", ".jpg", ".jpeg"}
OUTPUT_FOLDER = "Converted_MD_Files"
MAX_RETRIES = 3
WARN_TOKENS = 30_000  # avisa si un .md estimado supera esto

# Descripción de imágenes dentro de PDFs
MIN_IMAGE_SIDE = 150      # ignora imágenes más chicas (logos, íconos)
MAX_IMAGES_PER_PAGE = 3
MAX_IMAGES_PER_PDF = 30   # tope de gasto por archivo
MAX_CONSECUTIVE_FAILS = 3  # fallos seguidos de la API tras los cuales se abandona el PDF
CACHE_FILE = ".image_cache.json"  # descripciones ya pagadas, en la carpeta de salida
MAX_DESC_TOKENS = 400         # descripción breve; si se corta, queda marcado
REASONING_MARGIN = 4000   # extra para modelos que "piensan" (usan max_completion_tokens)

# Descripción breve, para dar contexto a la IA que lea el .md. No transcribe cifras
# (una lectura errónea pasa desapercibida); si hacen falta, se le pasa la imagen original.
IMAGE_PROMPT = (
    "Describí esta imagen en 1 a 3 oraciones, en texto plano y sin encabezados Markdown. "
    "Indicá qué tipo de imagen es (gráfico, tabla, esquema, plano, foto, captura...) y qué "
    "representa o para qué sirve en un documento técnico. Si tiene título, ejes, columnas o "
    "etiquetas, nombralos con sus unidades; en un gráfico, indicá además la forma general "
    "de la curva (creciente, con máximo, etc.). "
    "NO transcribas valores numéricos ni tablas y no estimes cifras. "
    "Si la imagen contiene datos numéricos, terminá con: Valores no transcriptos."
)



# ---------------------------------------------------------------- configuración
def get_env(name: str, default: str | None = None) -> str:
    value = os.environ.get(name, default)
    if not value:
        sys.exit(f"Falta la variable de entorno {name}. Ver instrucciones en el encabezado.")
    return value


LLM_VARS = ("AZURE_OPENAI_API_KEY", "AZURE_OPENAI_ENDPOINT", "AZURE_OPENAI_DEPLOYMENT")
DEFAULT_OPENAI_MODEL = "gpt-5.4-mini"  # con visión y barato; se cambia con OPENAI_MODEL en el .env


def build_converter():
    """Devuelve (MarkItDown, cliente, modelo). Proveedor de imágenes, en este orden:
    1) Azure OpenAI, si hay variables AZURE_OPENAI_*; 2) OpenAI, si hay OPENAI_API_KEY.
    Sin ninguna credencial el script funciona igual, pero sin describir imágenes.
    Si las de Azure están a medias, avisa cuál falta."""
    if any(os.environ.get(v) for v in LLM_VARS):
        client = AzureOpenAI(
            azure_endpoint=get_env("AZURE_OPENAI_ENDPOINT"),
            api_key=get_env("AZURE_OPENAI_API_KEY"),
            api_version=get_env("AZURE_OPENAI_API_VERSION", "2024-08-01-preview"),
        )
        model = get_env("AZURE_OPENAI_DEPLOYMENT")
        provider = "Azure OpenAI"
    elif os.environ.get("OPENAI_API_KEY"):
        client = OpenAI(api_key=os.environ["OPENAI_API_KEY"])  # también lee OPENAI_BASE_URL si existe
        model = os.environ.get("OPENAI_MODEL") or DEFAULT_OPENAI_MODEL
        provider = "OpenAI"
    else:
        print("Aviso: no hay credenciales de OpenAI ni de Azure OpenAI (archivo .env). "
              "Se convierte todo, pero sin describir imágenes.\n")
        return MarkItDown(), None, None
    print(f"Descripción de imágenes: {provider}, modelo {model}\n")
    md = MarkItDown(llm_client=client, llm_model=model, llm_prompt=IMAGE_PROMPT)
    return md, client, model


def clean_markdown(text: str) -> str:
    """Quita espacios sobrantes y líneas en blanco repetidas para ahorrar tokens."""
    text = text.replace("\r\n", "\n")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip() + "\n"


# ------------------------------------------------------------------- PDF: tablas
def _cell(value) -> str:
    return " ".join(str(value or "").split()).replace("|", "\\|")


def table_to_markdown(rows) -> str:
    """Convierte las filas de pdfplumber en una tabla Markdown."""
    rows = [[_cell(c) for c in r] for r in rows]
    rows = [r for r in rows if any(r)]
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]

    keep = [j for j in range(width) if any(r[j] for r in rows)]  # sin columnas vacías
    rows = [[r[j] for j in keep] for r in rows]
    width = len(keep)

    title = ""
    # Fila de título que abarca toda la tabla (una sola celda con texto)
    if width > 1 and len(rows) > 1 and sum(1 for c in rows[0] if c) == 1:
        title = f"**{next(c for c in rows[0] if c)}**\n\n"
        rows = rows[1:]

    header, body = rows[0], rows[1:]
    lines = ["| " + " | ".join(header) + " |", "|" + "---|" * width]
    lines += ["| " + " | ".join(r) + " |" for r in body]
    return title + "\n".join(lines)


def _fix_upright(page) -> None:
    """Algunos PDFs (tablas pegadas desde Excel/CAD) tienen el texto escalado dentro de un
    objeto y pdfplumber lo toma como vertical, separando cada letra. Se lo marca horizontal."""
    for c in page.chars:
        a, b, cc, d = c["matrix"][:4]
        if not c.get("upright", True) and a > 0 and d > 0 and abs(b) < 0.01 and abs(cc) < 0.01:
            c["upright"] = True


def _extract_text(page_or_crop) -> str:
    # y_tolerance mayor: mantiene subíndices (x_b, F_p) en la misma línea que su texto
    return (page_or_crop.extract_text(y_tolerance=5) or "").strip()


def _text_between(page, top: float, bottom: float) -> str:
    if bottom - top < 1:
        return ""
    x0, _, x1, _ = page.bbox
    return _extract_text(page.crop((x0, top, x1, bottom)))


def page_text_and_tables(page) -> str:
    """Texto de la página con las tablas convertidas a Markdown, en orden vertical."""
    _fix_upright(page)
    tables = []
    for t in page.find_tables():
        data = t.extract()
        if len(data) >= 2 and max(len(r) for r in data) >= 2:  # descarta cuadros de 1 fila/columna
            tables.append((t.bbox, data))
    tables.sort(key=lambda item: item[0][1])

    blocks, cursor = [], page.bbox[1]
    for (_, top, _, bottom), data in tables:
        if top < cursor:  # solapada con la anterior
            continue
        blocks.append(_text_between(page, cursor, top))
        blocks.append(table_to_markdown(data))
        cursor = bottom
    blocks.append(_text_between(page, cursor, page.bbox[3]))
    return "\n\n".join(b for b in blocks if b.strip())


# ------------------------------------------------------- PDF: imágenes con el LLM
_token_param = "max_tokens"  # cambia solo si el modelo lo rechaza (modelos nuevos)


def _chat_create(client, model, messages):
    """Llama al LLM. Los modelos nuevos (serie o*, GPT-5) rechazan max_tokens y piden
    max_completion_tokens, que además incluye los tokens de razonamiento."""
    global _token_param

    def call():
        n = MAX_DESC_TOKENS + (REASONING_MARGIN if _token_param == "max_completion_tokens" else 0)
        return client.chat.completions.create(model=model, messages=messages, **{_token_param: n})

    try:
        return call()
    except Exception as e:
        if _token_param == "max_tokens" and "max_completion_tokens" in str(e):
            _token_param = "max_completion_tokens"
            return call()
        raise


class ImageCache:
    """Caché en disco: (modelo + prompt + hash de la imagen) -> descripción.
    Evita volver a pagar la API por la misma imagen, por ejemplo al usar --force."""

    def __init__(self, path: Path):
        self.path = path
        try:
            self.data = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            self.data = {}

    @staticmethod
    def key(model: str, digest: str) -> str:
        prompt_id = hashlib.md5(IMAGE_PROMPT.encode("utf-8")).hexdigest()[:8]
        return f"{model}|{prompt_id}|{digest}"  # si cambia el modelo o el prompt, no se reutiliza

    def get(self, key: str):
        return self.data.get(key)

    def set(self, key: str, desc: str) -> None:
        self.data[key] = desc
        try:
            tmp = self.path.with_suffix(".tmp")
            tmp.write_text(json.dumps(self.data, ensure_ascii=False), encoding="utf-8")
            tmp.replace(self.path)
        except OSError as e:
            print(f"  Aviso: no se pudo guardar el caché de imágenes ({e.__class__.__name__}).")


def new_image_budget() -> dict:
    return {"left": MAX_IMAGES_PER_PDF, "fails": 0, "stopped": False}


def describe_page_images(client, model, cache, page, seen, budget):
    """Describe las imágenes raster de una página. `budget` lleva el estado del PDF:
    imágenes restantes, fallos seguidos de la API y si ya se abandonó."""
    out = []
    for img in page.images:
        if len(out) >= MAX_IMAGES_PER_PAGE or budget["left"] <= 0 or budget["stopped"]:
            break
        try:
            pil = img.image.convert("RGB")
            if min(pil.size) < MIN_IMAGE_SIDE:
                continue
            digest = hashlib.md5(img.data).hexdigest()
            if digest in seen:  # misma imagen repetida (logos, fondos)
                continue
            seen.add(digest)
        except Exception as e:  # imagen ilegible: no es un problema de la API
            print(f"  No se pudo leer una imagen ({e.__class__.__name__}): {e}")
            continue

        key = ImageCache.key(model, digest)
        desc = cache.get(key) if cache else None
        if desc is None:
            try:
                pil.thumbnail((1024, 1024))
                buf = io.BytesIO()
                pil.save(buf, format="JPEG", quality=85)
                b64 = base64.b64encode(buf.getvalue()).decode()
                r = _chat_create(client, model, [{"role": "user", "content": [
                    {"type": "text", "text": IMAGE_PROMPT},
                    {"type": "image_url",
                     "image_url": {"url": f"data:image/jpeg;base64,{b64}"}},
                ]}])
            except Exception as e:
                budget["fails"] += 1
                print(f"  No se pudo describir una imagen ({e.__class__.__name__}): {e}")
                if budget["fails"] >= MAX_CONSECUTIVE_FAILS:
                    budget["stopped"] = True
                    print(f"  La API falló {MAX_CONSECUTIVE_FAILS} veces seguidas: no se describen más "
                          "imágenes de este PDF. Revisá la clave, el endpoint y el nombre del modelo o deployment.\n"
                          "  Este .md queda sin esas descripciones y la próxima corrida lo salteará: "
                          "reconvertilo con --force cuando lo arregles (el caché evita pagar las ya hechas).")
                continue
            budget["fails"] = 0
            budget["left"] -= 1
            choice = r.choices[0]
            desc = (choice.message.content or "").strip()
            if not desc:
                if choice.finish_reason == "length":
                    print("  Aviso: una descripción salió vacía (el modelo agotó el límite pensando).")
                continue
            if choice.finish_reason == "length":
                desc += "\n(DESCRIPCIÓN TRUNCADA por el límite de tokens)"
                print("  Aviso: una descripción se cortó por el límite de tokens.")
            elif cache:
                cache.set(key, desc)  # las truncadas no se guardan: conviene reintentarlas
        else:
            budget["left"] -= 1  # el caché cuenta para el tope, así el resultado no cambia entre corridas
        out.append(f"[Imagen {len(out) + 1}: {desc}]")
    return out


EDGE_LINES = 3  # cuántas líneas del borde superior/inferior de cada página se revisan


def _edge_indices(n_lines: int) -> set:
    return set(range(min(EDGE_LINES, n_lines))) | set(range(max(0, n_lines - EDGE_LINES), n_lines))


def _is_page_number(key: str) -> bool:
    """¿La línea (con los dígitos reemplazados por #) es solo un número de página?
    Ej.: '#', '# / #', 'Página # de #', '- # -'. Filas de datos con varios números no."""
    k = key.lower()
    if k.count("#") > 2 or len(k) > 25:
        return False
    return re.sub(r"#|\W|_|p[áa]gina|p[áa]g|page|de|of", "", k) == ""


def strip_repeated_lines(pages: list) -> list:
    """Quita encabezados y pies de página que se repiten en la mayoría de las páginas.
    Un encabezado idéntico en todas se conserva solo en la primera página donde aparece;
    los números de página que varían ('12 / 16', '13 / 16') se quitan de todas.
    Las líneas que solo cambian en los dígitos pero no son números de página (por ejemplo
    filas de datos) no se tocan. Tampoco se tocan las tablas."""
    n = len(pages)
    if n < 2:
        return pages
    key = lambda line: re.sub(r"\d+", "#", line.strip())
    is_candidate = lambda line: bool(line.strip()) and not line.lstrip().startswith("|")

    split = [text.split("\n") for text in pages]
    seen_at = {}  # clave -> {n.º de página: línea original}
    for pi, lines in enumerate(split):
        for li in _edge_indices(len(lines)):
            if is_candidate(lines[li]):
                seen_at.setdefault(key(lines[li]), {})[pi] = lines[li].strip()

    need = max(2, -(-n * 6 // 10))  # al menos el 60% de las páginas (mínimo 2)
    repeated = {
        k: v for k, v in seen_at.items()
        if len(v) >= need and (len(set(v.values())) == 1 or _is_page_number(k))
    }

    out = []
    for pi, lines in enumerate(split):
        edge, kept = _edge_indices(len(lines)), []
        for li, line in enumerate(lines):
            k = key(line)
            if li in edge and is_candidate(line) and k in repeated:
                identical = len(set(repeated[k].values())) == 1
                if identical and pi == min(repeated[k]):
                    kept.append(line)  # el encabezado se conserva una vez
                continue
            kept.append(line)
        out.append("\n".join(kept).strip())
    return out


def convert_pdf_by_pages(path: Path, describe=None) -> str:
    """PDF -> Markdown con marcador por página, tablas separadas por celda, sin encabezados
    ni pies repetidos y, opcionalmente, descripción de imágenes.
    `describe` = (client, model, caché) o None."""
    reader = PdfReader(str(path)) if (describe and PdfReader) else None
    seen, budget = set(), new_image_budget()
    texts = []
    with pdfplumber.open(str(path)) as pdf:
        for i, page in enumerate(pdf.pages, 1):
            try:
                texts.append(page_text_and_tables(page))
            except Exception as e:
                print(f"  Página {i}: falló la extracción de tablas ({e.__class__.__name__}), uso texto simple.")
                texts.append(_extract_text(page))
    total = len(texts)
    texts = strip_repeated_lines(texts)

    parts = [f"<!-- {path.name} | {total} páginas -->"]
    for i, text in enumerate(texts, 1):
        flag = f" | {FORMULA_MARK}" if MATH_RE.search(text) else ""
        if reader is not None:
            descs = describe_page_images(*describe, reader.pages[i - 1], seen, budget)
            if descs:
                text = (text + "\n\n" + "\n".join(descs)).strip()
        parts.append(f"<!-- Página {i}/{total}{flag} -->\n{text}")
    return "\n\n".join(parts)


# ------------------------------------------------------------------ PPTX: notas
def strip_pptx_notes(text: str) -> str:
    """Quita los bloques '### Notes:' (notas del orador) que MarkItDown agrega por diapositiva."""
    return re.sub(r"\n#{1,6} Notes:\n.*?(?=\n<!-- Slide number:|\Z)", "", text, flags=re.S)


# ------------------------------------------------------------ XLSX: por bloques
XLSX_MAX_ROWS = 200  # filas máximas por bloque; el resto se omite con un aviso


def _fmt_number(v) -> str:
    """Número sin redondeos que cambien el valor (0.00000785 no pasa a 0.000008).
    Usa 15 cifras significativas, las que guarda Excel: no se pierde información real
    y se evita el ruido de coma flotante (0.1 + 0.2 sale 0.3, no 0.30000000000000004)."""
    if isinstance(v, int):
        return str(v)
    if v != v or v in (float("inf"), float("-inf")):
        return str(v)
    if v == int(v) and abs(v) < 1e15:
        return str(int(v))
    return format(Decimal(f"{v:.15g}"), "f")


def _xlsx_text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "VERDADERO" if v else "FALSO"
    if isinstance(v, (int, float)):
        return _fmt_number(v)
    if isinstance(v, datetime):
        return v.date().isoformat() if v.time() == datetime.min.time() else v.isoformat(sep=" ")
    if isinstance(v, date):
        return v.isoformat()
    return " ".join(str(v).split()).replace("|", "\\|")


def _xlsx_blocks(ws):
    """Agrupa las celdas con contenido en bloques (tablas separadas por filas/columnas vacías).
    Las celdas combinadas en vertical repiten su valor en cada fila."""
    grid = {}
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None and str(c.value).strip() != "":
                grid[(c.row, c.column)] = c.value
    occupied = set(grid)
    for mr in ws.merged_cells.ranges:
        top_left = grid.get((mr.min_row, mr.min_col))
        if top_left is None:
            continue
        vertical = mr.min_col == mr.max_col and mr.max_row > mr.min_row
        for r in range(mr.min_row, mr.max_row + 1):
            for c in range(mr.min_col, mr.max_col + 1):
                occupied.add((r, c))
                if vertical:
                    grid[(r, c)] = top_left

    seen, blocks = set(), []
    for start in sorted(occupied):
        if start in seen:
            continue
        comp, stack = [], [start]
        seen.add(start)
        while stack:
            r, c = stack.pop()
            comp.append((r, c))
            for nb in ((r + 1, c), (r - 1, c), (r, c + 1), (r, c - 1)):
                if nb in occupied and nb not in seen:
                    seen.add(nb)
                    stack.append(nb)
        blocks.append(comp)
    return grid, blocks


def _render_block(grid, comp) -> str:
    cells = set(comp)
    rows = sorted({r for r, c in comp if (r, c) in grid})
    cols = sorted({c for r, c in comp if (r, c) in grid})
    if not rows:
        return ""
    rng = (f"{get_column_letter(min(c for _, c in comp))}{min(r for r, _ in comp)}:"
           f"{get_column_letter(max(c for _, c in comp))}{max(r for r, _ in comp)}")
    matrix = [[_xlsx_text(grid[(r, c)]) if (r, c) in cells and (r, c) in grid else "" for c in cols]
              for r in rows]
    is_text = [[isinstance(grid.get((r, c)), str) and (r, c) in cells for c in cols] for r in rows]

    if sum(1 for row in matrix for x in row if x) == 1:  # celda suelta
        return f"{next(x for row in matrix for x in row if x)} ({rng})"

    title = ""
    if len(rows) > 1 and sum(1 for x in matrix[0] if x) == 1:  # fila de título
        title = next(x for x in matrix[0] if x)
        matrix, is_text = matrix[1:], is_text[1:]
    head = f"**{title}** ({rng})" if title else f"Rango {rng}"

    first = [t for x, t in zip(matrix[0], is_text[0]) if x]
    has_header = len(matrix) > 1 and len(first) >= 2 and all(first)
    body = matrix[1:] if has_header else matrix
    omitted = max(0, len(body) - XLSX_MAX_ROWS)
    body = body[:XLSX_MAX_ROWS]

    lines = [head]
    if has_header:
        lines.append("| " + " | ".join(matrix[0]) + " |")
        lines.append("|" + "---|" * len(cols))
        lines += ["| " + " | ".join(r) + " |" for r in body]
    else:
        for r in body:
            r = list(r)
            while r and not r[-1]:
                r.pop()
            lines.append(" | ".join(r))
    if omitted:
        lines.append(f"(... {omitted} filas más no incluidas)")
    return "\n".join(lines)


def convert_xlsx(path: Path) -> str:
    """Excel -> Markdown: una tabla por bloque de datos (sin NaN ni encabezados falsos)."""
    wb = openpyxl.load_workbook(str(path), data_only=True)
    parts = []
    try:
        for ws in wb.worksheets:
            if ws.sheet_state != "visible":
                continue
            grid, blocks = _xlsx_blocks(ws)
            blocks.sort(key=lambda b: (min(r for r, _ in b), min(c for _, c in b)))
            rendered = [x for x in (_render_block(grid, b) for b in blocks) if x]
            parts.append(f"## {ws.title}\n\n" + ("\n\n".join(rendered) if rendered else "(hoja vacía)"))
    finally:
        wb.close()
    return "\n\n".join(parts)


# ---------------------------------------------------------------- conversión
def convert_source(md: MarkItDown, path: Path, describe=None, skip_notes=False) -> str:
    ext = path.suffix.lower()
    if ext == ".pdf" and pdfplumber is not None:
        try:
            return convert_pdf_by_pages(path, describe)
        except Exception as e:
            print(f"  No se pudo procesar por páginas ({e.__class__.__name__}), conversión normal.")
    if ext == ".xlsx" and openpyxl is not None:
        try:
            return convert_xlsx(path)
        except Exception as e:
            print(f"  No se pudo separar por bloques ({e.__class__.__name__}), conversión normal.")
    text = md.convert(str(path)).text_content
    if ext == ".pptx" and skip_notes:
        text = strip_pptx_notes(text)
    if ext == ".docx":
        text = re.sub(r"</?u>", "", text)  # subrayado en HTML: solo suma tokens
    return text


def convert_with_retries(md, path: Path, describe=None, skip_notes=False) -> str:
    for attempt in range(1, MAX_RETRIES + 1):
        try:
            return convert_source(md, path, describe, skip_notes)
        except Exception as e:
            if attempt == MAX_RETRIES:
                raise
            wait = 2 ** attempt
            print(f"  Error ({e.__class__.__name__}), reintento {attempt}/{MAX_RETRIES - 1} en {wait}s...")
            time.sleep(wait)


# -------------------------------------------------------------- tokens y benchmark
_encoding = False  # False = todavía no se probó; None = tiktoken no disponible


def _get_encoding(warn_missing: bool = False):
    """tiktoken (opcional: pip install tiktoken). La primera vez baja el vocabulario,
    así que necesita internet; si falla, se usa la estimación caracteres / 4 y se avisa por qué."""
    global _encoding
    if _encoding is False:
        try:
            import tiktoken
            _encoding = tiktoken.get_encoding("o200k_base")
        except ImportError:
            _encoding = None
            if warn_missing:
                print("Aviso: tiktoken no está instalado en este Python (probá: py -m pip install tiktoken). "
                      "Los tokens se estiman con caracteres / 4.\n")
        except Exception as e:  # instalado, pero no pudo bajar o cargar el vocabulario
            _encoding = None
            print(f"Aviso: tiktoken no pudo cargar su vocabulario ({e.__class__.__name__}: {e}). "
                  "Los tokens se estiman con caracteres / 4.\n")
    return _encoding


def count_tokens(text: str) -> int:
    enc = _get_encoding()
    if enc is not None:
        return len(enc.encode(text, disallowed_special=()))
    return len(text) // 4


def token_method() -> str:
    return "tiktoken (o200k_base)" if _get_encoding() is not None else "estimación (caracteres / 4)"


def benchmark_file(plain: MarkItDown, path: Path, text: str):
    """Compara contra la conversión estándar de MarkItDown, sin LLM ni ajustes.
    Devuelve (tokens_markitdown, tokens_este_script, tokens_de_descripciones) o None."""
    try:
        raw = count_tokens(plain.convert(str(path)).text_content)
    except Exception as e:
        print(f"  Benchmark: MarkItDown no pudo convertir el archivo ({e.__class__.__name__}).")
        return None
    if raw <= 0:
        return None
    imgs = sum(count_tokens(x) for x in IMG_RE.findall(text))
    return raw, count_tokens(text), imgs


def write_benchmark(rows, output_dir: Path) -> Path:
    """rows = [(nombre, tokens_markitdown, tokens_este_script)]. Escribe una tabla lista para el README."""
    lines = ["| Archivo | MarkItDown puro | MD_Converter | Cambio |", "|---|---:|---:|---:|"]
    for name, raw, ours in rows:
        lines.append(f"| {name} | {raw:,} | {ours:,} | {(ours - raw) / raw:+.0%} |")
    if len(rows) > 1:
        raw_t, ours_t = sum(r[1] for r in rows), sum(r[2] for r in rows)
        lines.append(f"| **Total** | **{raw_t:,}** | **{ours_t:,}** | **{(ours_t - raw_t) / raw_t:+.0%}** |")
    lines += ["", f"Tokens contados con {token_method()}. \"MarkItDown puro\" es la conversión estándar, "
                  "sin descripción de imágenes. \"MD_Converter\" incluye, si las hay, las descripciones "
                  "de imágenes y las notas para la IA."]
    out = output_dir / "benchmark.md"
    out.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return out


# ------------------------------------------------------------------ estadísticas
IMAGE_EXTS = {".png", ".jpg", ".jpeg"}
NOTE_PREFIX = "> Nota del conversor:"
IMAGE_NOTE_TEXT = (
    "las imágenes de este archivo están descriptas sin transcribir valores numéricos. "
    "Antes de responder cualquier cifra que dependa de una imagen (un valor de una tabla o de "
    "un gráfico), pedile al usuario la imagen original indicando dónde está (página o "
    "diapositiva y número de imagen); no la estimes ni la reemplaces por cifras parecidas del texto."
)
FORMULA_MARK = "contiene fórmulas"
FORMULA_NOTE_TEXT = (
    'las páginas marcadas "contiene fórmulas" pueden tener las ecuaciones desordenadas '
    "(fracciones y subíndices separados del resto). Si una fórmula importa para la respuesta, "
    "pedile al usuario esa página del original."
)
# Letras matemáticas de Unicode (ecuaciones de Word), barra de fracción y símbolos ∑ ∫ √ ∏ ∂
MATH_RE = re.compile("[\U0001D400-\U0001D7FF\u2044\u2211\u222B\u221A\u220F\u2202]")
PAGE_RE = re.compile(r"<!-- Página (\d+)/(\d+)(?: \|[^>]*)? -->")
IMG_RE = re.compile(r"\[Imagen(?: \d+)?:.*?\](?=\n\[Imagen|\n\n<!-- Página|\s*\Z)", re.S)


def add_notes(text: str, path: Path, llm_enabled: bool) -> str:
    """Agrega al inicio una nota para la IA lectora: imágenes descriptas y/o páginas con fórmulas."""
    ext = path.suffix.lower()
    notes = []
    if llm_enabled and (
        ext in IMAGE_EXTS
        or (ext == ".pdf" and IMG_RE.search(text))
        or (ext == ".pptx" and re.search(r"!\[[^\]]+\]\(", text))
    ):
        notes.append(IMAGE_NOTE_TEXT)
    if ext == ".pdf" and FORMULA_MARK in text:
        notes.append(FORMULA_NOTE_TEXT)
    if not notes:
        return text
    head = f"{NOTE_PREFIX} {notes[0]}" + "".join(f"\n> Además, {n}" for n in notes[1:])
    return f"{head}\n\n{text}"


def section_stats(text: str):
    """[(título, tokens_estimados)] por cada encabezado de Markdown."""
    rows, current, buf = [], "(inicio)", []
    for line in text.splitlines():
        if line.startswith("#"):
            rows.append((current, count_tokens("\n".join(buf))))
            current, buf = line.lstrip("#").strip(), []
        else:
            buf.append(line)
    rows.append((current, count_tokens("\n".join(buf))))
    return [(t, n) for t, n in rows if n > 0]


def print_stats(text: str, top: int = 5, show_breakdown: bool = False):
    """Siempre informa las descripciones de imágenes (cuestan API). El detalle de páginas o
    secciones más pesadas solo se muestra si `show_breakdown` (--stats o archivo muy grande)."""
    imgs = IMG_RE.findall(text)
    if imgs:
        print(f"  Descripciones de imágenes: {len(imgs)} (~{sum(count_tokens(x) for x in imgs):,} tokens)")
    if not show_breakdown:
        return

    if PAGE_RE.search(text):  # PDF: pesa más por página que por encabezado
        chunks = PAGE_RE.split(text)  # [previo, n, total, texto, n, total, texto, ...]
        pages = [(int(chunks[i]), count_tokens(chunks[i + 2])) for i in range(1, len(chunks), 3)]
        if len(pages) > 1:
            print("  Páginas más pesadas:")
            for n, tokens in sorted(pages, key=lambda x: x[1], reverse=True)[:top]:
                print(f"    ~{tokens:>7,} tokens  página {n}")
        return

    stats = section_stats(text)
    if len(stats) >= 2:
        print(f"  Secciones más pesadas (de {len(stats)}):")
        for title, tokens in sorted(stats, key=lambda x: x[1], reverse=True)[:top]:
            print(f"    ~{tokens:>7,} tokens  {title[:60]}")


# ------------------------------------------------------------------------- CLI
def needs_conversion(src: Path, dst: Path, force: bool) -> bool:
    if force or not dst.exists():
        return True
    return src.stat().st_mtime > dst.stat().st_mtime  # el original se modificó


def find_files(input_dir: Path, recursive: bool, output_dir: Path):
    pattern = input_dir.rglob("*") if recursive else input_dir.glob("*")
    for p in sorted(pattern):
        if not p.is_file() or p.suffix.lower() not in EXTENSIONS:
            continue
        if output_dir in p.parents:  # no procesar la carpeta de salida
            continue
        yield p


def ask_yes_no(question: str, default: bool = False) -> bool:
    """Pregunta Y/N por terminal. Si no hay terminal interactiva, usa el valor por defecto."""
    if not sys.stdin.isatty():
        return default
    hint = "[Y/n]" if default else "[y/N]"
    while True:
        answer = input(f"{question} {hint}: ").strip().lower()
        if not answer:
            return default
        if answer in ("y", "yes", "s", "si", "sí"):
            return True
        if answer in ("n", "no"):
            return False
        print("  Respondé Y o N.")


def main():
    parser = argparse.ArgumentParser(description="Convierte archivos a Markdown.")
    parser.add_argument("-i", "--input", default=".", help="carpeta de entrada")
    parser.add_argument("-o", "--output", default=OUTPUT_FOLDER, help="carpeta de salida")
    parser.add_argument("-r", "--recursive", action="store_true", help="incluir subcarpetas")
    parser.add_argument("--force", action="store_true", help="reconvertir aunque ya exista")
    parser.add_argument("--describe-pdf-images", action="store_true",
                        help="describir imágenes de PDFs sin preguntar (usa la API)")
    parser.add_argument("--no-describe-pdf-images", action="store_true",
                        help="no describir imágenes de PDFs y no preguntar")
    parser.add_argument("--stats", action="store_true",
                        help="mostrar siempre las páginas/secciones más pesadas "
                             "(por defecto solo si el archivo supera WARN_TOKENS)")
    parser.add_argument("--benchmark", action="store_true",
                        help="comparar los tokens de cada archivo contra MarkItDown puro y guardar benchmark.md "
                             "(incluye los ya convertidos, sin volver a convertirlos)")
    parser.add_argument("--skip-notes", action="store_true",
                        help="descartar las notas del orador de los PPTX sin preguntar")
    parser.add_argument("--keep-notes", action="store_true",
                        help="conservar las notas del orador de los PPTX y no preguntar")
    args = parser.parse_args()

    input_dir = Path(args.input).resolve()
    output_dir = Path(args.output).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    md, client, model = build_converter()
    llm_enabled = client is not None
    _get_encoding(warn_missing=args.benchmark)  # si falla tiktoken, avisa al principio
    files = list(find_files(input_dir, args.recursive, output_dir))
    if not files:
        print(f"No se encontraron archivos soportados ({', '.join(sorted(EXTENSIONS))}).")
        return

    print(f"Archivos encontrados: {len(files)}\n")

    def dst_for(src: Path) -> Path:
        return output_dir / src.relative_to(input_dir).parent / f"{src.name}.md"

    pending = [f for f in files if needs_conversion(f, dst_for(f), args.force)]
    pending_pdfs = [f for f in pending if f.suffix.lower() == ".pdf"]
    pending_pptx = [f for f in pending if f.suffix.lower() == ".pptx"]

    if pending_pdfs and pdfplumber is None:
        print("Aviso: falta pdfplumber (pip install pdfplumber). Los PDFs se convertirán sin "
              "marcadores de página ni separación de tablas.\n")

    # --- PDFs: ¿describir imágenes?
    can_describe = llm_enabled and pdfplumber is not None and PdfReader is not None
    if args.describe_pdf_images and not can_describe:
        print("Aviso: --describe-pdf-images se ignora (falta configurar OpenAI/Azure OpenAI, o instalar pdfplumber o pypdf).\n")
    if args.no_describe_pdf_images or not pending_pdfs or not can_describe:
        describe_images = False
    elif args.describe_pdf_images:
        describe_images = True
    else:
        describe_images = ask_yes_no(
            f"Hay {len(pending_pdfs)} PDF(s) para convertir. "
            "¿Agregar una descripción de sus imágenes con IA? (usa la API y suma tokens)"
        )
        print()
    describe = (client, model, ImageCache(output_dir / CACHE_FILE)) if describe_images else None

    # --- PPTX: ¿descartar notas del orador?
    if args.keep_notes or not pending_pptx:
        skip_notes = False
    elif args.skip_notes:
        skip_notes = True
    else:
        skip_notes = ask_yes_no(
            f"Hay {len(pending_pptx)} PPTX para convertir. "
            "¿Descartar las notas del orador? (ahorra tokens, pero se pierden)"
        )
        print()

    ok = skipped = failed = 0
    plain = MarkItDown() if args.benchmark else None
    bench_rows = []

    def run_benchmark(rel, src, text):
        if plain is None or src.suffix.lower() in IMAGE_EXTS:
            return
        result = benchmark_file(plain, src, text)
        if result is None:
            return
        raw, ours, imgs = result
        bench_rows.append((str(rel), raw, ours))
        msg = f"  Benchmark: MarkItDown puro ~{raw:,} -> este script ~{ours:,} tokens ({(ours - raw) / raw:+.0%})"
        if imgs:
            msg += f"; incluye ~{imgs:,} de descripciones de imágenes"
        print(msg)

    for src in files:
        # Se conserva la extensión original: informe.pdf -> informe.pdf.md
        rel = src.relative_to(input_dir)
        dst = dst_for(src)
        dst.parent.mkdir(parents=True, exist_ok=True)

        if not needs_conversion(src, dst, args.force):
            print(f"Salteado: {rel} (ya convertido)")
            skipped += 1
            if plain is not None:
                run_benchmark(rel, src, dst.read_text(encoding="utf-8"))
            continue

        if not llm_enabled and src.suffix.lower() in IMAGE_EXTS:
            print(f"Salteado: {rel} (describir imágenes requiere configurar OpenAI o Azure OpenAI en el .env)")
            skipped += 1
            continue

        print(f"Convirtiendo: {rel}")
        try:
            text = clean_markdown(add_notes(convert_with_retries(md, src, describe, skip_notes), src, llm_enabled))
            dst.write_text(text, encoding="utf-8")
            est_tokens = count_tokens(text)  # tiktoken si está instalado; si no, caracteres / 4
            msg = f"  OK (~{est_tokens:,} tokens)"
            if est_tokens > WARN_TOKENS:
                msg += "  <-- muy grande, considerá recortarlo"
            print(msg)
            print_stats(text, show_breakdown=args.stats or est_tokens > WARN_TOKENS)
            run_benchmark(rel, src, text)
            ok += 1
        except Exception as e:
            print(f"  Error: {e}")
            failed += 1

    print(f"\nListo. Convertidos: {ok} | Salteados: {skipped} | Con error: {failed}")
    if bench_rows:
        out = write_benchmark(bench_rows, output_dir)
        raw_t, ours_t = sum(r[1] for r in bench_rows), sum(r[2] for r in bench_rows)
        print(f"Benchmark [{token_method()}]: {raw_t:,} -> {ours_t:,} tokens ({(ours_t - raw_t) / raw_t:+.0%}). "
              f"Tabla guardada en {out}")


if __name__ == "__main__":
    main()